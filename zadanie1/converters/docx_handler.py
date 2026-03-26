"""Conversion of Excel sheet data to DOCX (Word) format."""

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# Mapping of alignment names to python-docx constants
ALIGN_MAP = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
}

# Mapping of Excel alignment names to python-docx constants
EXCEL_ALIGN_MAP = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "general": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "centre": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
}

# Default column width in Excel units
DEFAULT_EXCEL_WIDTH = 8.43


def _parse_rgb(value) -> RGBColor | None:
    """
    Converts hex color string to RGBColor object.
    :param value: Color value to parse; expected to be a string like "#RRGGBB" or "#AARRGGBB".
    :return: RGBColor object or None if parsing fails.
    """
    if not isinstance(value, str):
        return None
    raw = value.strip().lstrip("#")
    if len(raw) == 8:
        raw = raw[2:]  # Strip alpha channel
    if len(raw) != 6:
        return None
    try:
        return RGBColor(int(raw[0:2], 16), int(raw[2:4], 16), int(raw[4:6], 16))
    except ValueError:
        return None


def _get_para_alignment(h_align: str | None, fallback) -> WD_ALIGN_PARAGRAPH:
    """
    Resolves paragraph alignment from an Excel horizontal-alignment name.
    :param h_align: Horizontal alignment string from Excel metadata.
    :param fallback: Alignment constant to use when h_align is absent or unknown.
    :return: Resolved WD_ALIGN_PARAGRAPH constant.
    """
    if isinstance(h_align, str):
        return EXCEL_ALIGN_MAP.get(h_align.lower().strip(), fallback)
    return fallback


def _normalize_widths(excel_widths: list, n_cols: int) -> list[float]:
    """
    Fills missing column widths with the default Excel column width.
    :param excel_widths: Raw column width values from the Excel file.
    :param n_cols: Expected number of columns.
    :return: List of n_cols float width values.
    """
    widths = []
    for w in (excel_widths or [])[:n_cols]:
        widths.append(float(w) if w not in (None, "") else DEFAULT_EXCEL_WIDTH)
    while len(widths) < n_cols:
        widths.append(DEFAULT_EXCEL_WIDTH)
    return widths


def _scale_widths_to_twips(
    excel_widths: list[float], available_twips: int
) -> list[int]:
    """
    Scales Excel column widths proportionally to fit the available page width.
    :param excel_widths: Column widths in Excel units.
    :param available_twips: Total available page width in twips.
    :return: List of column widths in twips, each at least 1.
    """
    total = sum(excel_widths) or 1
    return [max(int(available_twips * w / total), 1) for w in excel_widths]


def _get_font_scale(max_cols: int) -> float:
    """
    Chooses a font scale factor based on the number of columns.
    :param max_cols: Number of columns in the widest sheet.
    :return: Scale factor to apply to the original font size.
    """
    if max_cols <= 10:
        return 1.0
    if max_cols <= 15:
        return 0.85
    if max_cols <= 25:
        return 0.72
    return 0.6


def _add_page_numbers(footer) -> None:
    """
    Inserts an automatic page-number field into the document footer.
    :param footer: The footer object of a document section.
    :return: None
    """
    para = footer.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


def _shade_cell(cell, fill_hex: str) -> None:
    """
    Applies a background fill color to a table cell.
    :param cell: The python-docx table cell to shade.
    :param fill_hex: Hex color string (with or without leading "#").
    :return: None
    """
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex.lstrip("#").upper())
    existing = tcPr.find(qn("w:shd"))
    if existing is not None:
        tcPr.remove(existing)
    tcPr.append(shd)


def _set_cell_width(cell, width_twips: int) -> None:
    """
    Sets the explicit width of a table cell in twips.
    :param cell: The python-docx table cell to resize.
    :param width_twips: Desired cell width in twips.
    :return: None
    """
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcW = OxmlElement("w:tcW")
    tcW.set(qn("w:w"), str(width_twips))
    tcW.set(qn("w:type"), "dxa")
    existing = tcPr.find(qn("w:tcW"))
    if existing is not None:
        tcPr.remove(existing)
    tcPr.append(tcW)


def _apply_cell_formatting(
    cell,
    style: dict,
    font_scale: float,
    line_spacing: float,
    space_after: int,
    default_alignment,
) -> None:
    """
    Applies font, color, and paragraph formatting to a table cell.
    :param cell: The python-docx table cell to format.
    :param style: Style metadata with optional font, fill, and alignment sub-dicts.
    :param font_scale: Multiplier applied to the original font size.
    :param line_spacing: Line spacing value for all paragraphs in the cell.
    :param space_after: Space after each paragraph in points.
    :param default_alignment: Fallback paragraph alignment constant.
    :return: None
    """
    font_info = style.get("font") or {}
    fill_info = style.get("fill") or {}
    align_info = style.get("alignment") or {}

    # Apply background fill color
    if fill_info.get("color"):
        _shade_cell(cell, fill_info["color"])

    # Scale font size, enforcing a minimum of 6 pt
    original_size = float(font_info.get("size") or 11.0)
    scaled_size = max(original_size * font_scale, 6.0)

    for para in cell.paragraphs:
        para.alignment = _get_para_alignment(
            align_info.get("horizontal"), default_alignment
        )
        para.paragraph_format.line_spacing = line_spacing
        para.paragraph_format.space_after = Pt(space_after)

        for run in para.runs:
            run.font.size = Pt(scaled_size)
            if font_info.get("bold"):
                run.bold = True
            if font_info.get("italic"):
                run.italic = True
            if font_info.get("underline"):
                run.underline = True
            color = _parse_rgb(font_info.get("color"))
            if color:
                run.font.color.rgb = color


def _get_cell_meta(cells: list, row: int, col: int) -> dict:
    """
    Safely retrieves cell metadata from the nested cells list.
    :param cells: Nested list of cell metadata dicts from the xlsx reader.
    :param row: Zero-based row index.
    :param col: Zero-based column index.
    :return: Cell metadata dict, or an empty dict if the index is out of range.
    """
    try:
        return cells[row][col]
    except (IndexError, TypeError):
        return {}


def _setup_section(section, n_cols: int) -> int:
    """
    Configures page size, orientation, and margins for a document section.
    :param section: The python-docx section object to configure.
    :param n_cols: Number of columns in the sheet; determines orientation.
    :return: Available content width in twips after subtracting margins.
    """
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    if n_cols > 6:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = Cm(29.7), Cm(21.0)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width, section.page_height = Cm(21.0), Cm(29.7)

    margin_twips = int(Cm(2.0).pt * 20) * 2
    page_w_twips = int(section.page_width.pt * 20)
    return page_w_twips - margin_twips


def _add_sheet_table(
    doc: Document,
    df: pd.DataFrame,
    cells: list,
    merges: list,
    twip_widths: list[int],
    font_scale: float,
    line_spacing: float,
    space_after: int,
    default_alignment,
) -> None:
    """
    Adds a formatted table for a single sheet to the document.
    :param doc: Target python-docx Document object.
    :param df: Sheet data; each cell value is written to the table.
    :param cells: Nested list of cell metadata dicts (style, merge info, etc.).
    :param merges: List of merge-range dicts describing merged cell regions.
    :param twip_widths: Column widths in twips.
    :param font_scale: Font size multiplier.
    :param line_spacing: Line spacing for all paragraphs.
    :param space_after: Space after each paragraph in points.
    :param default_alignment: Default paragraph alignment constant.
    :return: None
    """
    n_rows, n_cols = df.shape
    table = doc.add_table(rows=n_rows, cols=n_cols)
    table.style = "Table Grid"

    # Merge cells according to Excel merge regions
    for merge in merges:
        if merge.get("is_anchor"):
            try:
                table.cell(merge["start_row"], merge["start_col"]).merge(
                    table.cell(merge["end_row"], merge["end_col"])
                )
            except (IndexError, ValueError):
                pass

    # Populate cell values and apply formatting
    for row_i in range(n_rows):
        for col_j in range(n_cols):
            meta = _get_cell_meta(cells, row_i, col_j)

            if meta.get("hidden_by_merge"):
                continue

            cell = table.cell(row_i, col_j)
            value = (
                meta.get("value")
                if meta.get("value") is not None
                else str(df.iat[row_i, col_j])
            )
            cell.text = str(value)

            _set_cell_width(cell, twip_widths[col_j])
            _apply_cell_formatting(
                cell,
                meta.get("style", {}),
                font_scale,
                line_spacing,
                space_after,
                default_alignment,
            )

    doc.add_paragraph()


def df_to_docx(sheets: dict, settings: dict, output) -> None:
    """
    Converts Excel sheet data to a Word document and writes it to output.
    :param sheets: Mapping of sheet names to payload dicts as returned by read_xlsx().
    Each payload must contain a dataframe key and may contain excel_col_widths, cells, and merges.
    :param settings: Formatting settings with keys: alignment, line_spacing, space_after,
    page_numbers, and title.
    :param output: Writable binary stream (e.g. BytesIO) that receives the generated .docx content.
    :return: None
    """
    alignment = ALIGN_MAP.get(
        settings.get("alignment", "left"), WD_ALIGN_PARAGRAPH.LEFT
    )
    line_spacing = float(settings.get("line_spacing", 1.15))
    space_after = int(settings.get("space_after", 6))
    page_numbers = settings.get("page_numbers", True)
    title = settings.get("title", "").strip()

    # Prepare and validate sheet data
    prepared = {}
    for name, payload in sheets.items():
        df = payload.get("dataframe")
        if df is None or not isinstance(df, pd.DataFrame) or df.empty:
            continue
        df = df.fillna("")
        prepared[name] = {
            "df": df,
            "widths": _normalize_widths(
                payload.get("excel_col_widths", []), df.shape[1]
            ),
            "cells": payload.get("cells", []),
            "merges": payload.get("merges", []),
        }

    if not prepared:
        return

    # Build the document
    doc = Document()

    if title:
        heading = doc.add_heading(title, level=0)
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Each sheet gets its own section (new page)
    for idx, (sheet_name, data) in enumerate(prepared.items()):
        n_cols = data["df"].shape[1]
        font_scale = _get_font_scale(n_cols)

        if idx == 0:
            section = doc.sections[0]
            available_twips = _setup_section(section, n_cols)
        else:
            doc.add_section()
            section = doc.sections[-1]
            available_twips = _setup_section(section, n_cols)

        if len(prepared) > 1:
            doc.add_heading(sheet_name, level=1).alignment = alignment

        twip_widths = _scale_widths_to_twips(data["widths"], available_twips)
        _add_sheet_table(
            doc,
            data["df"],
            data["cells"],
            data["merges"],
            twip_widths,
            font_scale,
            line_spacing,
            space_after,
            alignment,
        )

    if page_numbers:
        for section in doc.sections:
            _add_page_numbers(section.footer)

    doc.save(output)
