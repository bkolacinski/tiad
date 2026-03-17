import io
from copy import copy

import openpyxl
import pandas as pd
from openpyxl.utils import get_column_letter

DEFAULT_COLUMN_WIDTH = 8.43
DEFAULT_FONT_SIZE = 11.0

# Indexed colors based on Excel's default palette (0-23).
_INDEXED_COLORS = {
    0: "#000000",
    1: "#FFFFFF",
    2: "#FF0000",
    3: "#00FF00",
    4: "#0000FF",
    5: "#FFFF00",
    6: "#FF00FF",
    7: "#00FFFF",
    8: "#000000",
    9: "#FFFFFF",
    10: "#FF0000",
    11: "#00FF00",
    12: "#0000FF",
    13: "#FFFF00",
    14: "#FF00FF",
    15: "#00FFFF",
    16: "#800000",
    17: "#008000",
    18: "#000080",
    19: "#808000",
    20: "#800080",
    21: "#008080",
    22: "#C0C0C0",
    23: "#808080",
}


def _normalize_color(value) -> str | None:
    """
    Converts Openpyxl color to hex RGB format.
    :param value: Openpyxl color object or None.
    :return: "#RRGGBB" or None if color is not valid.
    """
    if not value:
        return None

    rgb = getattr(value, "rgb", None)
    if rgb and isinstance(rgb, str):
        rgb = rgb.upper()
        if len(rgb) == 8:
            rgb = rgb[2:]  # Skip alfa
        if len(rgb) == 6:
            return f"#{rgb}"

    indexed = getattr(value, "indexed", None)
    if indexed is not None:
        try:
            return _INDEXED_COLORS.get(indexed)
        except (TypeError, KeyError):
            return None

    return None


def _extract_fill_color(fill) -> str | None:
    """
    Extracts the fill color. Only solid and gray125 patterns are considered.
    :param fill: Openpyxl fill object.
    :return: Hex color string or None if not applicable.
    """
    patterns = {"solid", "gray125"}
    if not fill or getattr(fill, "patternType", None) not in patterns:
        return None
    return _normalize_color(getattr(fill, "fgColor", None)) or _normalize_color(
        getattr(fill, "start_color", None)
    )


def _extract_border(border) -> dict:
    """
    Extracts border information for all sides of the cell.
    Only sides with defined styles are included.
    :param border: Cell border object from Openpyxl.
    :return: Dictionary with border styles and colors for each side
    (left, right, top, bottom). Empty if no borders are defined.
    """
    if not border:
        return {}
    result = {}
    for side_name in ("left", "right", "top", "bottom"):
        side = getattr(border, side_name, None)
        if side and getattr(side, "style", None):
            result[side_name] = {
                "style": side.style,
                "color": _normalize_color(getattr(side, "color", None)),
            }
    return result


def _extract_alignment(alignment) -> dict:
    """
    Extracts text alignment information (horizontal, vertical, wrap_text).
    :param alignment: Openpyxl alignment object.
    :return: Dictionary with alignment properties.
    Empty if alignment is not defined.
    """
    if not alignment:
        return {}
    return {
        "horizontal": alignment.horizontal,
        "vertical": alignment.vertical,
        "wrap_text": bool(alignment.wrap_text),
    }


def _extract_font(font) -> dict:
    """
    Returns Openpyxl font information.
    :param font: Openpyxl font object.
    :return: Dictionary with font information.
    """
    if not font:
        return {}
    return {
        "name": font.name,
        "size": float(font.sz) if font.sz is not None else DEFAULT_FONT_SIZE,
        "bold": bool(font.b),
        "italic": bool(font.i),
        "underline": bool(font.u),
        "color": _normalize_color(getattr(font, "color", None)),
    }


def _extract_cell_style(cell) -> dict:
    """
    Returns Openpyxl cell style information.
    :param cell: Openpyxl cell object.
    :return: Dictionary with cell style information.
    """
    return {
        "font": _extract_font(cell.font),
        "fill": {"color": _extract_fill_color(cell.fill)},
        "border": _extract_border(cell.border),
        "alignment": _extract_alignment(cell.alignment),
        "number_format": cell.number_format or "General",
    }


def _collect_merged_map(ws, n_rows: int, n_cols: int) -> dict[tuple[int, int], dict]:
    """
    Returns a map of merged cells
    :param ws: Openpyxl worksheet object.
    :param n_rows: Number of rows.
    :param n_cols: Number of columns.
    :return: Map of merged cells.
    """
    merged_map: dict[tuple[int, int], dict] = {}

    for merged_range in ws.merged_cells.ranges:
        min_col, min_row = merged_range.min_col, merged_range.min_row
        max_col = min(merged_range.max_col, n_cols)
        max_row = min(merged_range.max_row, n_rows)

        if min_row > n_rows or min_col > n_cols:
            continue

        meta = {
            "range": str(merged_range),
            "start_row": min_row - 1,
            "start_col": min_col - 1,
            "end_row": max_row - 1,
            "end_col": max_col - 1,
            "is_anchor": True,
        }

        merged_map[(min_row, min_col)] = meta

        for row in range(min_row, max_row + 1):
            for col in range(min_col, max_col + 1):
                if row == min_row and col == min_col:
                    continue
                hidden_meta = meta.copy()
                hidden_meta["is_anchor"] = False
                merged_map[(row, col)] = hidden_meta

    return merged_map


def _collect_column_widths(ws, n_cols: int) -> list[float]:
    """
    Returns a list of column widths.
    :param ws: Openpyxl worksheet object.
    :param n_cols: Number of columns.
    :return: List of column widths.
    """
    widths = []
    for i in range(1, n_cols + 1):
        dim = ws.column_dimensions.get(get_column_letter(i))
        widths.append(
            float(dim.width or DEFAULT_COLUMN_WIDTH) if dim else DEFAULT_COLUMN_WIDTH
        )
    return widths


def _build_dataframe(ws, n_rows: int, n_cols: int) -> pd.DataFrame:
    """
    Creates a DataFrame with cell values from the worksheet.
    :param ws: Openpyxl worksheet object.
    :param n_rows: Number of rows.
    :param n_cols: Number of columns.
    :return: pandas DataFrame with cell values.
    """
    rows = []
    for row_idx in range(1, n_rows + 1):
        row = []
        for col_idx in range(1, n_cols + 1):
            val = ws.cell(row=row_idx, column=col_idx).value
            row.append(str(val) if val is not None else "")
        rows.append(row)
    return pd.DataFrame(rows, dtype=str).fillna("")


def _collect_cells_metadata(
        ws, n_rows: int, n_cols: int, merged_map: dict[tuple[int, int], dict]
) -> list[list[dict]]:
    """
    Collects cell metadata from the worksheet.
    :param ws: Openpyxl worksheet object.
    :param n_rows: Number of rows.
    :param n_cols: Number of columns.
    :param merged_map: Map of merged cells.
    :return: List of cell metadata.
    """
    cells_meta = []

    for row_idx in range(1, n_rows + 1):
        row_meta = []
        for col_idx in range(1, n_cols + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            merge_meta = merged_map.get((row_idx, col_idx))

            if merge_meta and not merge_meta["is_anchor"]:
                row_meta.append(
                    {
                        "value": "",
                        "style": {},
                        "merge": merge_meta,
                        "hidden_by_merge": True,
                    }
                )
                continue

            val = str(cell.value) if cell.value is not None else ""
            row_meta.append(
                {
                    "value": val,
                    "style": _extract_cell_style(cell),
                    "merge": merge_meta,
                    "hidden_by_merge": False,
                }
            )
        cells_meta.append(row_meta)

    return cells_meta


def _find_used_range(ws) -> tuple[int, int]:
    """
    Finds data range, skips empty rows/cols at the end of worksheet.
    :param ws: Openpyxl worksheet object.
    :return: Tuple with last row and last col which contain data.
    """
    max_row_raw = ws.max_row or 0
    max_col_raw = ws.max_column or 0

    if max_row_raw == 0 or max_col_raw == 0:
        return 0, 0

    last_col = 0
    for row in ws.iter_rows(
            min_row=1, max_row=max_row_raw, min_col=1, max_col=max_col_raw
    ):
        for cell in reversed(row):
            if cell.value is not None:
                last_col = max(last_col, cell.column)
                break

    last_row = 0
    for row in ws.iter_rows(
            min_row=1, max_row=max_row_raw, min_col=1, max_col=max(last_col, 1)
    ):
        for cell in row:
            if cell.value is not None:
                last_row = max(last_row, cell.row)
                break

    return last_row, last_col


def read_xlsx(file) -> dict[str, dict]:
    """
    Reads data from a xlsx file.
    :param file: File to read.
    :return: Dictionary with all data read:
    sheet_name -> {dataframe, excel_col_widths, cells, merges}
    """
    wb = openpyxl.load_workbook(io.BytesIO(file.read()), data_only=True)
    result = {}

    for name in wb.sheetnames:
        ws = wb[name]
        n_rows, n_cols = _find_used_range(ws)

        if n_rows == 0 or n_cols == 0:
            continue

        merged_map = _collect_merged_map(ws, n_rows, n_cols)
        df = _build_dataframe(ws, n_rows, n_cols)
        cells_meta = _collect_cells_metadata(ws, n_rows, n_cols, merged_map)
        col_widths = _collect_column_widths(ws, n_cols)

        # Get unique merges (only anchors) to avoid duplicates in output.
        seen = set()
        merges = []
        for meta in merged_map.values():
            if meta["is_anchor"] and meta["range"] not in seen:
                seen.add(meta["range"])
                merges.append(copy(meta))

        result[name] = {
            "dataframe": df,
            "excel_col_widths": col_widths,
            "cells": cells_meta,
            "merges": merges,
        }

    return result
